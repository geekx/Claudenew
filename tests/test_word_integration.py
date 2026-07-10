"""
Word 端到端集成测试（需真实 python-docx）。
验证：正文段落 + 表格单元 + 文本框 全部被翻译、零漏译；
文本框提取覆盖 DrawingML / VML / mc:AlternateContent 去重。

若未安装 python-docx 则跳过。运行：python tests/test_word_integration.py
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from conftest import load_ncat

try:
    import docx  # noqa
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn
    HAVE_DOCX = True
except Exception:
    HAVE_DOCX = False

NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
      'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
      'xmlns:v="urn:schemas-microsoft-com:vml" '
      'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"')


def _drawingml_textbox(text):
    return (f'<w:r {NS}><w:drawing><wp:inline><wp:extent cx="1" cy="1"/><wp:docPr id="1" name="t"/>'
            f'<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            f'<wps:wsp><wps:txbx><w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p>'
            f'</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')


def test_textbox_extraction_variants():
    ncat = load_ncat(use_real_docx=True)
    # DrawingML 文本框
    doc = Document()
    p = doc.add_paragraph(); p.add_run()._r.addnext(parse_xml(_drawingml_textbox('文本框文字')))
    els = []
    assert ncat.extract_docx_textbox_paragraphs(doc, els) == 1
    assert [t for _, t in els] == ['文本框文字']
    # mc:AlternateContent：Choice + Fallback 只取一次
    doc2 = Document()
    alt = (f'<w:r {NS}><mc:AlternateContent>'
           f'<mc:Choice Requires="wps"><w:drawing><wp:inline><wp:extent cx="1" cy="1"/><wp:docPr id="1" name="t"/>'
           f'<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
           f'<wps:wsp><wps:txbx><w:txbxContent><w:p><w:r><w:t>选择版</w:t></w:r></w:p></w:txbxContent></wps:txbx></wps:wsp>'
           f'</a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice>'
           f'<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>回退版</w:t></w:r></w:p>'
           f'</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>')
    p2 = doc2.add_paragraph(); p2.add_run()._r.addnext(parse_xml(alt))
    els2 = []
    assert ncat.extract_docx_textbox_paragraphs(doc2, els2) == 1
    assert [t for _, t in els2] == ['选择版']
    # 无文本框
    doc3 = Document(); doc3.add_paragraph('普通')
    assert ncat.extract_docx_textbox_paragraphs(doc3, []) == 0
    print('test_textbox_extraction_variants: PASS')


def test_word_end_to_end_no_drop():
    ncat = load_ncat(use_real_docx=True)
    ncat._translation_memory = ncat.TranslationMemory(os.path.join(tempfile.mkdtemp(), 'tm.db'))

    doc = Document()
    doc.add_paragraph('第一段：网络安全策略概述')
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = '表格单元甲'
    t.rows[0].cells[1].text = '表格单元乙'
    p = doc.add_paragraph(); p.add_run()._r.addnext(parse_xml(_drawingml_textbox('文本框中的重要提示')))
    path_in = os.path.join(tempfile.mkdtemp(), 'in.docx')
    doc.save(path_in)

    doc2 = Document(path_in)
    elements = []

    def grab(c):
        if c is None:
            return
        for para in getattr(c, 'paragraphs', []):
            if para.text.strip():
                elements.append((para, para.text))
        for tb in getattr(c, 'tables', []):
            for row in tb.rows:
                for cell in row.cells:
                    grab(cell)
    grab(doc2)
    ncat.extract_docx_textbox_paragraphs(doc2, elements)
    assert '文本框中的重要提示' in [t for _, t in elements]

    class FakeJA:
        def translate_batch(self, texts, s, t, formality='auto', context_hint=''):
            return ['やく〔' + x + '〕です' for x in texts]
        def _translate_single(self, text, s, t, formality='auto', context_hint=''):
            return 'やく〔' + text + '〕です'

    class CB:
        def __init__(self): self.errors = []
        def progress(self, m): pass
        def status(self, m): pass
        def info(self, m): pass
        def error(self, m): self.errors.append(m)

    cb = CB()
    ncat.process_text_elements(elements, {}, None, FakeJA(), 'Chinese', 'Japanese', cb,
                               'replace', 'auto',
                               {'confidential': False, 'pivot': False, 'contextInference': False, 'fuzzyTM': True})
    path_out = os.path.join(tempfile.mkdtemp(), 'out.docx')
    doc2.save(path_out)

    out = Document(path_out)
    all_text = []

    def collect(c):
        if c is None:
            return
        for para in getattr(c, 'paragraphs', []):
            if para.text.strip():
                all_text.append(para.text)
        for tb in getattr(c, 'tables', []):
            for row in tb.rows:
                for cell in row.cells:
                    collect(cell)
    collect(out)
    el2 = []
    ncat.extract_docx_textbox_paragraphs(out, el2)
    all_text += [t for _, t in el2]

    untranslated = [t for t in all_text if 'やく〔' not in t]
    assert not untranslated, f'漏译: {untranslated}'
    assert cb.errors == [], f'不应有漏译报警: {cb.errors}'
    print('test_word_end_to_end_no_drop: PASS')


def test_tm_end_to_end_reuse():
    """翻译记忆端到端复用：同一句二次翻译应零 API 调用（精确+模糊命中）。"""
    ncat = load_ncat(use_real_docx=True)
    ncat._translation_memory = ncat.TranslationMemory(os.path.join(tempfile.mkdtemp(), 'tm.db'))

    class CountingJA:
        def __init__(self): self.translated = []
        def translate_batch(self, texts, s, t, formality='auto', context_hint=''):
            self.translated += list(texts)
            return ['やく〔' + x + '〕です' for x in texts]
        def _translate_single(self, text, s, t, formality='auto', context_hint=''):
            self.translated.append(text)
            return 'やく〔' + text + '〕です'

    class CB:
        def __init__(self): self.errors = []; self.infos = []
        def progress(self, m): pass
        def status(self, m): pass
        def info(self, m): self.infos.append(m)
        def error(self, m): self.errors.append(m)

    def run(text, translator):
        d = Document(); d.add_paragraph(text)
        els = [(p, p.text) for p in d.paragraphs if p.text.strip()]
        cb = CB()
        ncat.process_text_elements(els, {}, None, translator, 'Chinese', 'Japanese', cb,
                                   'replace', 'auto',
                                   {'confidential': False, 'pivot': False, 'contextInference': False, 'fuzzyTM': True})
        return d.paragraphs[0].text, cb

    sent = '网络安全是企业的重中之重'
    ft1 = CountingJA(); out1, _ = run(sent, ft1)
    assert out1.startswith('やく〔') and len(ft1.translated) == 1
    ft2 = CountingJA(); out2, cb2 = run(sent, ft2)
    assert out2 == out1 and len(ft2.translated) == 0          # 精确命中，零API
    assert any('TM' in i for i in cb2.infos)
    ft3 = CountingJA(); _, _ = run(sent + '。', ft3)
    assert len(ft3.translated) == 0                            # 模糊命中，零API
    print('test_tm_end_to_end_reuse: PASS')


def test_dominant_run_formatting():
    """多run段落的译文应套用「携带字符最多」的run格式，而非第一个run。
    段首短加粗标签不应让整段译文变粗；反之主导run加粗时译文应加粗。"""
    ncat = load_ncat(use_real_docx=True)

    # 场景1：短加粗前缀 + 长正常正文 → 译文不应加粗
    d = Document()
    p = d.add_paragraph()
    r1 = p.add_run('重要：'); r1.bold = True
    p.add_run('这是一段很长的正文说明内容，占据段落绝大部分字符数量。')
    ncat.update_element_text(p, 'やく〔本文〕です')
    assert len(p.runs) == 1
    assert p.runs[0].bold is not True, '短加粗前缀不应让整段译文加粗'

    # 场景2：短正常前缀 + 长加粗正文 → 译文应加粗
    d2 = Document()
    p2 = d2.add_paragraph()
    p2.add_run('注：')
    r2 = p2.add_run('这一大段加粗文字才是段落的主体内容所在位置。'); r2.bold = True
    ncat.update_element_text(p2, 'やく〔本文〕です')
    assert p2.runs[0].bold is True, '主导run加粗时译文应加粗'
    print('test_dominant_run_formatting: PASS')


def test_single_pass_analysis():
    """process_text_elements 应只做一遍术语保护（曾经分析+回填各算一遍）。"""
    ncat = load_ncat(use_real_docx=True)
    ncat._translation_memory = ncat.TranslationMemory(os.path.join(tempfile.mkdtemp(), 'tm.db'))

    calls = []
    orig = ncat.protect_and_replace_terms
    def counting(text, term_map, term_re):
        calls.append(text)
        return orig(text, term_map, term_re)
    ncat.protect_and_replace_terms = counting
    try:
        d = Document()
        for i in range(4):
            d.add_paragraph(f'第{i}段中文内容需要翻译处理')
        els = [(p, p.text) for p in d.paragraphs if p.text.strip()]

        class FakeJA:
            def translate_batch(self, texts, s, t, formality='auto', context_hint=''):
                return ['やく〔' + x + '〕です' for x in texts]
            def _translate_single(self, text, s, t, formality='auto', context_hint=''):
                return 'やく〔' + text + '〕です'

        class CB:
            def progress(self, m): pass
            def status(self, m): pass
            def info(self, m): pass
            def error(self, m): pass

        ncat.process_text_elements(els, {}, None, FakeJA(), 'Chinese', 'Japanese', CB(),
                                   'replace', 'auto',
                                   {'confidential': False, 'pivot': False, 'contextInference': False, 'fuzzyTM': True})
    finally:
        ncat.protect_and_replace_terms = orig
    assert len(calls) == 4, f'应每元素保护一次，实际调用{len(calls)}次'
    assert all(p.text.startswith('やく〔') for p in d.paragraphs), '单遍化后翻译结果不应变化'
    print('test_single_pass_analysis: PASS')


if __name__ == '__main__':
    if not HAVE_DOCX:
        print('⚠ python-docx 未安装，跳过 Word 集成测试')
        sys.exit(0)
    test_textbox_extraction_variants()
    test_word_end_to_end_no_drop()
    test_tm_end_to_end_reuse()
    test_dominant_run_formatting()
    test_single_pass_analysis()
    print('\n✅ Word 集成测试全部通过')
