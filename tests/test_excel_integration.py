"""
Excel 端到端集成测试（需真实 openpyxl）。
验证：文本单元格被翻译；公式单元格（含内嵌中文的公式）保持原样不被送译、不被破坏。
若未安装 openpyxl 则跳过。运行：python tests/test_excel_integration.py
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from conftest import load_ncat

try:
    import openpyxl  # noqa
    HAVE_XL = True
except Exception:
    HAVE_XL = False


def test_is_formula_cell():
    ncat = load_ncat(use_real_openpyxl=True)
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws['A1'] = '网络安全报告'
    ws['A2'] = '=SUM(B1:B3)'
    ws['A3'] = '=CONCATENATE(A1," 附注")'
    assert ncat.is_formula_cell(ws['A2']) is True
    assert ncat.is_formula_cell(ws['A3']) is True
    assert ncat.is_formula_cell(ws['A1']) is False
    print('test_is_formula_cell: PASS')


def test_excel_skips_formulas_end_to_end():
    ncat = load_ncat(use_real_openpyxl=True)
    ncat._translation_memory = ncat.TranslationMemory(os.path.join(tempfile.mkdtemp(), 'tm.db'))
    import openpyxl

    wb = openpyxl.Workbook(); ws = wb.active
    ws['A1'] = '网络安全报告'                    # 文本 → 应翻译
    ws['A2'] = '机密数据汇总'                    # 文本 → 应翻译
    ws['B1'] = '=SUM(C1:C9)'                     # 公式 → 不译
    ws['B2'] = '=CONCATENATE(A1," 附注")'        # 公式含中文 → 不译，不得破坏
    path_in = os.path.join(tempfile.mkdtemp(), 'in.xlsx')
    wb.save(path_in)

    # 仿真翻译流程的 Excel 抽取（对齐 run_translation_process 的 xlsx 分支）
    wb2 = openpyxl.load_workbook(path_in)
    elements = []
    for sheet in wb2.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if (cell.value and isinstance(cell.value, str) and cell.value.strip()
                        and not ncat.is_formula_cell(cell)):
                    elements.append((cell, cell.value))
    extracted = sorted({t for _, t in elements})
    assert '网络安全报告' in extracted and '机密数据汇总' in extracted
    assert '=SUM(C1:C9)' not in extracted
    assert '=CONCATENATE(A1," 附注")' not in extracted

    class FakeJA:
        def translate_batch(self, texts, s, t, formality='auto', context_hint=''):
            return ['やく〔' + x + '〕です' for x in texts]
        def _translate_single(self, text, s, t, formality='auto', context_hint=''):
            return 'やく〔' + text + '〕です'

    class CB:
        def __init__(self): self.errors = []
        def progress(self, m): pass
        def status(self, m): pass
        def info(self, m): pass
        def error(self, m): self.errors.append(m)

    ncat.process_text_elements(elements, {}, None, FakeJA(), 'Chinese', 'Japanese', CB(),
                               'replace', 'auto',
                               {'confidential': False, 'pivot': False, 'contextInference': False, 'fuzzyTM': True})
    path_out = os.path.join(tempfile.mkdtemp(), 'out.xlsx')
    wb2.save(path_out)

    out = openpyxl.load_workbook(path_out)
    ws2 = out.active
    assert ws2['A1'].value.startswith('やく〔')       # 文本已翻译
    assert ws2['A2'].value.startswith('やく〔')
    assert ws2['B1'].value == '=SUM(C1:C9)'           # 公式原样保留
    assert ws2['B2'].value == '=CONCATENATE(A1," 附注")'
    print('test_excel_skips_formulas_end_to_end: PASS')


if __name__ == '__main__':
    if not HAVE_XL:
        print('⚠ openpyxl 未安装，跳过 Excel 集成测试')
        sys.exit(0)
    test_is_formula_cell()
    test_excel_skips_formulas_end_to_end()
    print('\n✅ Excel 集成测试通过')
